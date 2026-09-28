"""Test stampa PDF calendario (/calendario/stampa): elenco testuale, non griglia."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from io import BytesIO

from docx import Document

from lys_workflow_hub.config import Settings
from lys_workflow_hub.core.pratica_eventi_repository import PraticaEventiRepository
from lys_workflow_hub.core.utenti_repository import UtentiRepository
from lys_workflow_hub.core.wincar_repository import (
    Cliente, CompagniaCliente, Controparte, Pratica, Sinistro, Veicolo,
)
from lys_workflow_hub.main import app
from lys_workflow_hub.web.routes import _genera_docx_calendario_stampa, get_app_settings, get_repository
from tests.conftest import ADMIN_EMAIL, ADMIN_PASSWORD, login_as

OPERATORE_EMAIL = "operatore-cal@test.local"
OPERATORE_PASSWORD = "test-password-1234"


def _sample_pratica(numero: int) -> Pratica:
    return Pratica(
        numero=numero,
        data_creazione=datetime(2026, 5, 1, 0, 0),
        cliente=Cliente("ROSSI MARIO", None, None, None, None, None, None, None, None, None),
        veicolo=Veicolo("AB123CD", None, None, None),
        sinistro=Sinistro(None, None, None, None, None, None, None),
        controparte=Controparte(None, None, None, None, None, None, None, None),
        assicurazione_cliente=CompagniaCliente(None, None, None, None, None, None, None),
    )


def test_genera_docx_calendario_stampa_contenuto():
    per_giorno = {
        date(2026, 9, 15): [
            {"evento": _EventoFake(766, "Perizia"), "cliente": "Rossi Mario", "targa": "AB123CD"},
        ],
    }
    docx_bytes = _genera_docx_calendario_stampa(per_giorno, "Settembre", 2026)
    doc = Document(BytesIO(docx_bytes))
    testo = "\n".join(p.text for p in doc.paragraphs)
    assert "Settembre 2026" in testo
    assert "15/09/2026" in testo
    assert "Perizia" in testo
    assert "Rossi Mario" in testo
    assert "AB123CD" in testo
    assert "pratica 766" in testo


def test_genera_docx_calendario_stampa_mese_vuoto():
    docx_bytes = _genera_docx_calendario_stampa({}, "Settembre", 2026)
    doc = Document(BytesIO(docx_bytes))
    testo = "\n".join(p.text for p in doc.paragraphs)
    assert "Nessun appuntamento" in testo


class _EventoFake:
    def __init__(self, pratica_numero: int, titolo: str):
        self.pratica_numero = pratica_numero
        self.titolo = titolo


@pytest.fixture
def client_stampa(tmp_path: Path, authenticated_app: UtentiRepository):
    settings = Settings(app_db_path=tmp_path / "app.db", wincar_archivio=tmp_path)
    eventi_repo = PraticaEventiRepository(db_path=settings.app_db_path)
    eventi_repo.add(766, "Perizia danni", date(2026, 9, 15), creato_da=1, creato_da_nome="Admin")

    wincar_repo = MagicMock()
    wincar_repo.get_pratica.return_value = _sample_pratica(766)

    app.dependency_overrides[get_app_settings] = lambda: settings
    app.dependency_overrides[get_repository] = lambda: wincar_repo
    try:
        yield TestClient(app), settings
    finally:
        app.dependency_overrides.pop(get_app_settings, None)
        app.dependency_overrides.pop(get_repository, None)


def test_calendario_stampa_admin_ok(client_stampa):
    client, _ = client_stampa
    login_as(client, ADMIN_EMAIL, ADMIN_PASSWORD)
    with patch("lys_workflow_hub.web.routes.docx_bytes_to_pdf_bytes", return_value=b"%PDF-fake") as mock_pdf:
        resp = client.get("/calendario/stampa", params={"anno": 2026, "mese": 9})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert "appuntamenti_2026-09.pdf" in resp.headers["content-disposition"]
    mock_pdf.assert_called_once()


def test_calendario_stampa_negata_a_non_admin(client_stampa):
    client, _ = client_stampa
    app.state.utenti_repo.create(
        email=OPERATORE_EMAIL, password=OPERATORE_PASSWORD, nome="Operatore Test", ruolo="operatore"
    )
    login_as(client, OPERATORE_EMAIL, OPERATORE_PASSWORD)
    resp = client.get("/calendario/stampa", params={"anno": 2026, "mese": 9})
    assert resp.status_code == 403
